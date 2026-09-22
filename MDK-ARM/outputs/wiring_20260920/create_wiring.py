from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.page import PageMargins

OUT = Path(__file__).resolve().parent / 'STM32F407VET6_外设接线表.xlsx'
wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, headers, rows, widths, table_name):
    ws = wb.create_sheet(name)
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).font = Font(name='微软雅黑', size=18, bold=True, color='FFFFFF')
    ws.cell(1, 1).fill = PatternFill('solid', fgColor='17365D')
    ws.row_dimensions[1].height = 36
    ws.append(['依据当前工程源码整理 · 2026-09-20 · 引脚为 GPIO 名称，非 PCB 插座针脚编号'])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1).font = Font(name='微软雅黑', size=10, color='526779')
    ws.row_dimensions[2].height = 26
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.font = Font(name='微软雅黑', size=11, color='203449')
            c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[row[0].row].height = 48
    for c in ws[3]:
        c.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='236B8E')
    ws.row_dimensions[3].height = 28
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(3, i).column_letter].width = width
    table = Table(displayName=table_name, ref=f'A3:{ws.cell(ws.max_row, len(headers)).coordinate}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = 'C4'
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:3'
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:{ws.cell(ws.max_row, len(headers)).coordinate}'
    ws.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4, header=.15, footer=.15)
    ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
    return ws

serial = '115200 bit/s，8N1'
rows = [
 ['Qt / USB 串口','PA9','USART1_TX / AF7','USB 串口 RXD',serial,'已启用','主控发送遥测','Core/Src/usart.c'],
 ['Qt / USB 串口','PA10','USART1_RX / AF7','USB 串口 TXD',serial,'已启用','接收上位机命令','Core/Src/usart.c'],
 ['OPS 定位板','PD5','USART2_TX / AF7','OPS PA3 / USART2_RX',serial,'已启用','F407 → OPS','Core/Src/usart.c；OPS Core/Src/usart.c'],
 ['OPS 定位板','PD6','USART2_RX / AF7','OPS PA2 / USART2_TX',serial,'已启用','OPS → F407','Core/Src/usart.c；OPS Core/Src/usart.c'],
 ['ZDT 底盘电机','PA0','UART4_TX / AF8','电机通信接口 RX',serial,'已启用','Emm 协议；轮地址见电机地址页','Core/Src/usart.c'],
 ['ZDT 底盘电机','PA1','UART4_RX / AF8','电机通信接口 TX',serial,'已启用','接收电机回复；总线接法按驱动器接口规范','Core/Src/usart.c'],
 ['摄像头串口','PB10','USART3_TX / AF7','摄像头 RX',serial,'预留','已初始化串口，尚无应用协议','Core/Src/usart.c'],
 ['摄像头串口','PB11','USART3_RX / AF7','摄像头 TX',serial,'预留','尚无接收处理应用','Core/Src/usart.c'],
 ['DM / 28 / 35 电机 CAN','PB6','CAN2_TX / AF9','CAN 收发器 TXD','1 Mbit/s','已启用','经收发器连接 CANH / CANL','Core/Src/can.c'],
 ['DM / 28 / 35 电机 CAN','PB5','CAN2_RX / AF9','CAN 收发器 RXD','1 Mbit/s','已启用','不可直接接 CANH / CANL','Core/Src/can.c'],
]
sheet('通信接口','STM32F407VET6 · 通信接口接线', ['外设','主控引脚','功能 / 复用','连接到外设端','通信参数','状态','备注','源码依据'], rows, [25,13,25,30,25,13,42,43], 'Communication')

rows = []
for pin, signal in [('PB13','CLK / SCL'),('PC3','DIN / SDA'),('PE2','RES / RST'),('PE3','DC'),('PE4','CS')]:
    rows.append(['OLED 软件 SPI',pin,'GPIO 输出',signal,'已启用','128×64 屏；SCL/SDA 名称不表示 I²C','Hardware/OLED_SoftSPI.h'])
rows += [
 ['夹爪舵机','PE9','TIM1_CH1 / AF1','舵机信号输入','预留','PWM 尚未启动；当前定时器配置不是常规舵机 PWM','Core/Src/tim.c'],
 ['电机电源使能','PD0','VM_EN 输出','外部电源开关 EN','GPIO 已配置','高有效，默认低；无自动使能；不是电机供电端','Core/Src/gpio.c；Core/Inc/main.h'],
 ['相机补光','PD1','CAMERA_LIGHT_EN 输出','灯驱动 EN','GPIO 已配置','高有效，默认低；不可直接驱动大功率灯','Core/Src/gpio.c；Core/Inc/main.h'],
 ['启动按键 1','PD2','START_KEY1 输入','按键一端；另一端 GND','预留','内部上拉，按下为低；未实现启动动作及消抖','Core/Src/gpio.c'],
 ['启动按键 2','PD3','START_KEY2 输入','按键一端；另一端 GND','预留','内部上拉，按下为低；未实现启动动作及消抖','Core/Src/gpio.c'],
 ['板载通信灯','PB2','COMM_LED 输出','板载 LED 电路','GPIO 已配置','高有效，默认熄灭；无自动闪烁','Core/Src/gpio.c'],
]
sheet('OLED及控制','STM32F407VET6 · OLED 与控制接口', ['外设','主控引脚','功能','连接到外设端','状态','备注','源码依据'], rows, [23,13,30,32,19,57,42], 'ControlPins')

sheet('下载与供电','STM32F407VET6 · SWD 与供电接线', ['接口','主控端','连接到','说明'], [
 ['SWD','PA13','下载器 SWDIO','调试数据'],['SWD','PA14','下载器 SWCLK','调试时钟'],
 ['复位','NRST','下载器 RESET','用于复位及复位下连接'],['公共地','GND','下载器及外设 GND','通信模块需共地'],
 ['电压参考','3.3V','下载器 VTref','目标板电平参考；不等同于下载器供电输出'],
 ['OLED 供电','按板卡电源设计','OLED VCC / GND','VCC 按具体 OLED 模块规格；源码不能确定模块供电电压'],
 ['其他外设供电','按板卡电源设计','OPS / 摄像头 / 电机电源','按模块额定电压独立核对；GPIO 不作为电机电源'],
 ], [24,24,36,86], 'DebugPower')

sheet('电机地址','底盘轮序与 ZDT 地址', ['ZDT 地址','轮位置','观察方向','通信接口'], [
 [1,'左前','俯视，车头朝前','UART4：PA0 TX / PA1 RX'],
 [2,'右前','俯视，车头朝前','UART4：PA0 TX / PA1 RX'],
 [3,'左后','俯视，车头朝前','UART4：PA0 TX / PA1 RX'],
 [4,'右后','俯视，车头朝前','UART4：PA0 TX / PA1 RX'],
 ], [20,24,35,55], 'MotorAddresses')

sheet('接线说明','使用说明与核对范围', ['项目','说明'], [
 ['工程根目录',r'E:\STM32\STM32F407VET6_ILHC'],
 ['OPS 源码目录',r'E:\STM32\ILHC\ops9-main\code'],
 ['引脚表示方法','PA9 等均为 MCU GPIO 名称，不是开发板插座序号或芯片封装脚号；插座位置需对照实际板卡原理图。'],
 ['串口方向','TX/RX 以主控为参考：主控 TX 接外设 RX，主控 RX 接外设 TX。逻辑电平按 3.3V 系统匹配。'],
 ['CAN 版本','当前使用 CAN2，PB5=RX、PB6=TX。旧文档中的 CAN1 PA11/PA12 不适用于当前代码。'],
 ['CAN 接线','MCU PB6/PB5 接收发器 TXD/RXD；收发器 CANH/CANL 再接电机总线。'],
 ['OLED 接口','当前为软件 SPI，PB13=CLK、PC3=DIN、PE2=RES、PE3=DC、PE4=CS；不要按 I²C 两线屏连接。'],
 ['状态含义','已启用表示有固件使用路径，不代表已验证实物接线；预留表示管脚或外设已配置，但相应应用功能未完成。'],
 ['供电依据','源码不能确定外设 VCC 额定电压、PCB 插座针序及外部驱动电路，请以实际模块规格和板卡原理图为准。'],
 ['电机轮序依据','Hardware/mecanum_control.c；俯视车头朝前：1 左前、2 右前、3 左后、4 右后。'],
 ], [28,120], 'WiringNotes')

wb.save(OUT)
check = load_workbook(OUT)
assert len(check.worksheets) == 5
assert check['通信接口']['B12'].value == 'PB6'
assert check['通信接口']['B13'].value == 'PB5'
assert sum(s.max_row - 3 for s in check.worksheets) == 42
for s in check:
    assert s.freeze_panes == 'C4'
    assert len(s.tables) == 1
    for row in s:
        for c in row:
            assert c.data_type != 'e'
print(OUT)
print('Verified: 5 worksheets, 42 data rows, all tables and frozen headers present.')
